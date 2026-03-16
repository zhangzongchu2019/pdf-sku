"""导出 22 列标准 Excel（与参考格式一致）。"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

# 22 列标准表头
STANDARD_HEADERS = [
    "商品图片",       # A
    "规格图片",       # B
    "*商品名称/描述", # C
    "售价",          # D
    "货号",          # E
    "商品ID",        # F
    "标签",          # G
    "来源(仅自己可见)", # H
    "商品简称",       # I
    "商品规格",       # J
    "颜色",          # K
    "商品编码",       # L
    "商品条码",       # M
    "市场价",        # N
    "成本价",        # O
    "库存",          # P
    "重量(kg)",      # Q
    "体积(m³)",      # R
    "商品详情",       # S
    "排序",          # T
    "上架状态",       # U
    "备注",          # V
]

# Pipeline attributes → Excel 列索引 (0-based)
ATTR_TO_COL = {
    "product_name": 2,   # C: *商品名称/描述
    "price": 3,          # D: 售价
    "model_number": 4,   # E: 货号
    "tag": 6,            # G: 标签
    "source": 7,         # H: 来源
    "specs": 9,          # J: 商品规格
    "size": 9,           # J: 商品规格 (兼容旧字段)
    "color": 10,         # K: 颜色
}


def export_dataset(
    run_result: dict,
    output_path: Path,
    source_name: str = "",
    tag: str = "",
) -> None:
    """将 Pipeline 输出导出为 22 列标准 Excel。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "商品数据"

    # 写表头
    header_font = Font(bold=True)
    for col_idx, header in enumerate(STANDARD_HEADERS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font

    # 写 SKU 数据
    row_num = 2
    sku_counter = 0
    for page in run_result.get("pages", []):
        for sku in page.get("skus", []):
            if sku.get("validity", "valid") != "valid":
                continue

            attrs = sku.get("attributes", {})
            sku_counter += 1

            # 映射 attributes 到列
            for attr_key, col_idx in ATTR_TO_COL.items():
                val = attrs.get(attr_key, "")
                if val:
                    ws.cell(row=row_num, column=col_idx + 1, value=str(val))

            # 自动生成货号（如果 pipeline 没提取到）
            if not attrs.get("model_number"):
                prefix = source_name[:2] if source_name else "SK"
                ws.cell(row=row_num, column=5, value=f"{prefix}{sku_counter:03d}")

            # 标签和来源
            if tag:
                ws.cell(row=row_num, column=7, value=tag)
            elif attrs.get("tag"):
                ws.cell(row=row_num, column=7, value=str(attrs["tag"]))

            if source_name:
                ws.cell(row=row_num, column=8, value=source_name)

            # A: 商品图片路径
            image_path = sku.get("image_path", "")
            if image_path:
                ws.cell(row=row_num, column=1, value=image_path)

            row_num += 1

    # 列宽调整
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["E"].width = 15
    ws.column_dimensions["G"].width = 15
    ws.column_dimensions["H"].width = 20
    ws.column_dimensions["J"].width = 25

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"已导出 {sku_counter} 个 SKU → {output_path}")
